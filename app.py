from flask import Flask, request, jsonify
from docx import Document
import io

app = Flask(__name__)


@app.route("/health", methods=["GET"])
def health():
    return "docx-extractor service is up", 200


@app.route("/extract", methods=["POST"])
def extract():
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400

    uploaded = request.files["file"]

    try:
        doc = Document(io.BytesIO(uploaded.read()))

        parts = [p.text for p in doc.paragraphs]

        # Contracts often put key terms (notice period, salary, etc.) in tables —
        # include table cell text too, not just body paragraphs.
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)

        text = "\n".join(p for p in parts if p)
        return jsonify({"text": text})

    except Exception as e:
        return jsonify({"error": f"Could not parse .docx file: {e}"}), 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=3000)
